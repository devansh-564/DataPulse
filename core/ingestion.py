import os
import io
import json
import xml.etree.ElementTree as ET
import pandas as pd
import numpy as np

def detect_file_format_and_category(file_name: str, file_bytes: bytes = None):
    """
    Detects file extension, MIME category (Structured, Semi-Structured, Unstructured),
    and processing strategy.
    """
    ext = os.path.splitext(file_name)[1].lower()
    
    if ext in ['.csv', '.xlsx', '.xls']:
        category = "Structured"
        format_type = ext.replace('.', '').upper()
    elif ext in ['.json', '.xml']:
        category = "Semi-Structured"
        format_type = ext.replace('.', '').upper()
    elif ext in ['.txt', '.pdf', '.docx']:
        category = "Unstructured"
        format_type = ext.replace('.', '').upper()
    else:
        category = "Unknown"
        format_type = ext.replace('.', '').upper() if ext else "RAW"

    return {
        "file_name": file_name,
        "extension": ext,
        "format_type": format_type,
        "category": category,
        "processing_strategy": f"{category} Data Pipeline"
    }

def ingest_file(file_obj_or_path, file_name: str = None):
    """
    Unified Ingestion Layer:
    Loads file from path or file-like object and returns standardized metadata,
    structured DataFrame (if tabular/semi-structured), or raw text & document metadata (if unstructured).
    """
    if isinstance(file_obj_or_path, str):
        file_name = file_name or os.path.basename(file_obj_or_path)
        with open(file_obj_or_path, 'rb') as f:
            file_bytes = f.read()
    elif hasattr(file_obj_or_path, 'read'):
        file_bytes = file_obj_or_path.read()
        if hasattr(file_obj_or_path, 'name'):
            file_name = file_name or file_obj_or_path.name
        if hasattr(file_obj_or_path, 'seek'):
            file_obj_or_path.seek(0)
    else:
        raise ValueError("Unsupported input source. Must be filepath or file-like object.")

    info = detect_file_format_and_category(file_name, file_bytes)
    ext = info["extension"]

    result = {
        "info": info,
        "file_bytes": file_bytes,
        "file_size_bytes": len(file_bytes),
        "structured_df": None,
        "unstructured_text": None,
        "document_metadata": {},
        "ingestion_status": "Success",
        "error_message": None
    }

    try:
        # 1. STRUCTURED FILES
        if ext == '.csv':
            try:
                df = pd.read_csv(io.BytesIO(file_bytes))
            except Exception:
                df = pd.read_csv(io.BytesIO(file_bytes), encoding='latin1')
            result["structured_df"] = df
            result["info"]["num_rows"] = len(df)
            result["info"]["num_cols"] = len(df.columns)

        elif ext in ['.xlsx', '.xls']:
            df = pd.read_excel(io.BytesIO(file_bytes))
            result["structured_df"] = df
            result["info"]["num_rows"] = len(df)
            result["info"]["num_cols"] = len(df.columns)

        # 2. SEMI-STRUCTURED FILES
        elif ext == '.json':
            try:
                json_data = json.loads(file_bytes.decode('utf-8'))
                if isinstance(json_data, list):
                    df = pd.DataFrame(json_data)
                elif isinstance(json_data, dict):
                    # Try to find records list or normalize dict
                    records_key = next((k for k, v in json_data.items() if isinstance(v, list)), None)
                    if records_key:
                        df = pd.DataFrame(json_data[records_key])
                    else:
                        df = pd.json_normalize(json_data)
                else:
                    df = pd.DataFrame([{"content": str(json_data)}])
                
                result["structured_df"] = df
                result["info"]["num_rows"] = len(df)
                result["info"]["num_cols"] = len(df.columns)
            except Exception as je:
                result["unstructured_text"] = file_bytes.decode('utf-8', errors='ignore')
                result["info"]["category"] = "Unstructured"

        elif ext == '.xml':
            try:
                df = pd.read_xml(io.BytesIO(file_bytes))
                result["structured_df"] = df
                result["info"]["num_rows"] = len(df)
                result["info"]["num_cols"] = len(df.columns)
            except Exception:
                # Fallback XML parsing
                root = ET.fromstring(file_bytes)
                records = []
                for child in root:
                    row = {sub.tag: sub.text for sub in child}
                    if row:
                        records.append(row)
                if records:
                    df = pd.DataFrame(records)
                    result["structured_df"] = df
                    result["info"]["num_rows"] = len(df)
                    result["info"]["num_cols"] = len(df.columns)
                else:
                    result["unstructured_text"] = file_bytes.decode('utf-8', errors='ignore')

        # 3. UNSTRUCTURED FILES
        elif ext == '.txt':
            raw_text = file_bytes.decode('utf-8', errors='ignore')
            result["unstructured_text"] = raw_text
            result["document_metadata"] = {
                "length_chars": len(raw_text),
                "length_words": len(raw_text.split()),
                "num_lines": len(raw_text.splitlines())
            }

        elif ext == '.pdf':
            text = ""
            num_pages = 0
            try:
                import fitz # PyMuPDF
                doc = fitz.open(stream=file_bytes, filetype="pdf")
                num_pages = len(doc)
                for page in doc:
                    text += page.get_text() + "\n"
            except Exception as pdfe:
                # Fallback simple reader if PyMuPDF not active yet
                text = f"[PDF Parsing Notice: {pdfe}] " + file_bytes.decode('latin1', errors='ignore')
            
            result["unstructured_text"] = text
            result["document_metadata"] = {
                "num_pages": num_pages,
                "length_chars": len(text),
                "length_words": len(text.split()),
                "has_text": len(text.strip()) > 0
            }

        elif ext == '.docx':
            text = ""
            try:
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                paras = [p.text for p in doc.paragraphs if p.text.strip()]
                text = "\n".join(paras)
                
                # Extract tables inside DOCX if present
                docx_tables = []
                for t in doc.tables:
                    t_rows = []
                    for row in t.rows:
                        t_rows.append([cell.text.strip() for cell in row.cells])
                    if t_rows:
                        docx_tables.append(t_rows)
                result["document_metadata"]["tables_found"] = len(docx_tables)
                result["document_metadata"]["extracted_tables"] = docx_tables
            except Exception as docxe:
                text = file_bytes.decode('utf-8', errors='ignore')
            
            result["unstructured_text"] = text
            result["document_metadata"]["length_chars"] = len(text)
            result["document_metadata"]["length_words"] = len(text.split())

    except Exception as e:
        result["ingestion_status"] = "Error"
        result["error_message"] = str(e)

    return result
